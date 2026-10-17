#!/usr/bin/env python3
"""厚生労働省「薬価基準収載品目リスト」から js/drug-list.js を作る。

使い方:
  git fetch origin source-data
  mkdir -p /tmp/src && git archive origin/source-data | tar -x -C /tmp/src
  python3 scripts/build-drug-list.py /tmp/src

source-data ブランチは .github/workflows/fetch-drug-sources.yml が作る。
必要なライブラリ: openpyxl, pypdf
"""
import glob
import json
import os
import re
import sys
import unicodedata

import openpyxl
import pypdf

KUBUN = {"内用薬": "内", "注射薬": "注", "外用薬": "外", "歯科用薬剤": "歯"}
OUT = os.path.join(os.path.dirname(__file__), "..", "js", "drug-list.js")


def nfkc(s):
    return unicodedata.normalize("NFKC", str(s or "")).strip()


def read_classes(pdf_path):
    text = "\n".join(p.extract_text() for p in pypdf.PdfReader(pdf_path).pages)
    classes = {}
    for line in text.splitlines():
        m = re.match(r"^(\d{3})\s+(\S.*)$", line.strip())
        if m:
            # "鎮痛，鎮痒，収斂，消炎剤" -> "鎮痛・鎮痒・収斂・消炎剤" (commas would read as a list of classes)
            classes[m.group(1)] = nfkc(m.group(2)).replace(",", "・")
    return classes


def is_item_list(ws):
    header = next(ws.iter_rows(max_row=1, values_only=True))
    return header[:3] == ("区分", "薬価基準収載医薬品コード", "成分名")


def main(src):
    classes = read_classes(os.path.join(src, "class.pdf"))
    rows = {}
    for path in sorted(glob.glob(os.path.join(src, "xlsx", "*.xlsx"))):
        ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
        if not is_item_list(ws):
            continue
        header = next(ws.iter_rows(max_row=1, values_only=True))
        col = {name: i for i, name in enumerate(header) if name}
        for r in ws.iter_rows(min_row=2, values_only=True):
            code = nfkc(r[col["薬価基準収載医薬品コード"]])
            kubun = KUBUN.get(nfkc(r[col["区分"]]))
            if not re.match(r"^[0-9A-Z]{12}$", code) or not kubun:
                continue
            # The same product can appear in both the April list and the latest list; the later file wins.
            rows[code] = (
                kubun,
                nfkc(r[col["成分名"]]),
                nfkc(r[col["規格"]]),
                nfkc(r[col["品名"]]),
            )

    # 成分名 -> YJコード上7桁(PMDAの成分ページ) -> 品名 -> 規格
    by_name = {}
    for code, (kubun, name, spec, product) in sorted(rows.items()):
        groups = by_name.setdefault(name, {})
        g = groups.setdefault(code[:7], {"k": kubun, "p": {}})
        specs = g["p"].setdefault(product, [])
        if spec and spec not in specs:
            specs.append(spec)

    items = []
    for name in sorted(by_name):
        groups = [
            [yj7, g["k"], [[product, "、".join(specs)] for product, specs in g["p"].items()]]
            for yj7, g in by_name[name].items()
        ]
        items.append([name, groups])

    used = sorted({yj7[:3] for _, groups in items for yj7, _, _ in groups})
    source_url = open(os.path.join(src, "source.txt")).read().strip()
    raw = open(os.path.join(src, "page.html"), "rb").read()
    charset = re.search(rb'charset=["\']?([\w-]+)', raw)
    page = nfkc(raw.decode(charset.group(1).decode() if charset else "utf-8", errors="replace"))
    m = re.search(r"(令和\d+年\d+月\d+日適用)", page)
    data = {
        "source": {
            "label": "厚生労働省「薬価基準収載品目リスト」" + ("(" + m.group(1) + ")" if m else ""),
            "url": source_url,
        },
        "classes": {c: classes[c] for c in used if c in classes},
        "items": items,
    }
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("// 自動生成ファイル(scripts/build-drug-list.py)。直接編集しないでください。\n")
        f.write("// 出典: " + data["source"]["label"] + " " + source_url + "\n")
        f.write("// items: [成分名, [[YJコード上7桁, 区分(内/注/外/歯), [[品名, 規格]...]]...]]\n")
        f.write("window.DRUG_LIST = ")
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    products = sum(len(p) for _, groups in items for _, _, p in groups)
    print(f"{len(rows)} 品目 / {len(items)} 成分 / {products} 販売名 -> {os.path.relpath(OUT)} ({os.path.getsize(OUT):,} bytes)")


if __name__ == "__main__":
    main(sys.argv[1])
